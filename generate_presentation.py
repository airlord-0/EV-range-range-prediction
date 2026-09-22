"""
PowerPoint Presentation Generator for EV Driving Range Prediction
Generates a 16:9 widescreen presentation matching the requested structure.
No emojis used. Professional academic and engineering theme.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation(output_path: str, assets_dir: str):
    prs = Presentation()
    # Set 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_slide_layout = prs.slide_layouts[6]

    # Color Palette (Corporate / Academic)
    COLOR_PRIMARY = RGBColor(15, 30, 60)       # Deep Navy
    COLOR_ACCENT = RGBColor(29, 78, 216)       # Blue
    COLOR_DARK_TEXT = RGBColor(30, 41, 59)     # Dark Slate
    COLOR_MUTED_TEXT = RGBColor(100, 116, 139) # Muted Gray
    COLOR_LIGHT_BG = RGBColor(248, 250, 252)   # Light Gray
    COLOR_BORDER = RGBColor(226, 232, 240)     # Border Slate
    COLOR_WHITE = RGBColor(255, 255, 255)

    def add_header(slide, title_text: str, category_text: str = ""):
        # Top banner background line
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.1))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = COLOR_ACCENT
        top_bar.line.fill.background()

        # Category tag
        if category_text:
            cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
            cat_tf = cat_box.text_frame
            cat_tf.word_wrap = True
            cat_p = cat_tf.paragraphs[0]
            cat_p.text = category_text.upper()
            cat_p.font.size = Pt(11)
            cat_p.font.bold = True
            cat_p.font.color.rgb = COLOR_ACCENT

        # Slide Title
        t_top = Inches(0.7) if category_text else Inches(0.5)
        title_box = slide.shapes.add_textbox(Inches(0.8), t_top, Inches(11.7), Inches(0.7))
        title_tf = title_box.text_frame
        title_tf.word_wrap = True
        title_p = title_tf.paragraphs[0]
        title_p.text = title_text
        title_p.font.size = Pt(24)
        title_p.font.bold = True
        title_p.font.color.rgb = COLOR_PRIMARY

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (Dark Theme)
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_slide_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_PRIMARY
    bg1.line.fill.background()

    # Blue accent line on title
    accent1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.8), Inches(1.2), Inches(0.08))
    accent1.fill.solid()
    accent1.fill.fore_color.rgb = COLOR_ACCENT
    accent1.line.fill.background()

    t_box1 = s1.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.5), Inches(3.2))
    tf1 = t_box1.text_frame
    tf1.word_wrap = True

    p1 = tf1.paragraphs[0]
    p1.text = "Electric Vehicle Driving Range Prediction"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE

    p2 = tf1.add_paragraph()
    p2.text = "Physics-Grounded Machine Learning Regression and Performance Benchmarking"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(191, 219, 254) # Light Blue
    p2.space_before = Pt(12)

    p3 = tf1.add_paragraph()
    p3.text = "Target Vehicle: 2013 Nissan Leaf (24 kWh Battery) | Reference: University of Michigan VED Schema"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(148, 163, 184) # Slate Muted
    p3.space_before = Pt(24)

    # -------------------------------------------------------------
    # SLIDE 2: 1. Introduction
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_slide_layout)
    add_header(s2, "1. Introduction", "Background and Motivation")

    # Left card: Context
    box2_l = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    box2_l.fill.solid()
    box2_l.fill.fore_color.rgb = COLOR_LIGHT_BG
    box2_l.line.color.rgb = COLOR_BORDER
    tf2_l = box2_l.text_frame
    tf2_l.word_wrap = True
    tf2_l.margin_left = Inches(0.3)
    tf2_l.margin_right = Inches(0.3)
    tf2_l.margin_top = Inches(0.3)

    p = tf2_l.paragraphs[0]
    p.text = "Industry Context and Transition"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    items_l = [
        "Global automotive shift toward Battery Electric Vehicles (BEVs) to achieve net-zero transport emissions.",
        "Range anxiety remains one of the primary deterrents to widespread consumer adoption.",
        "Unlike internal combustion vehicles with 500+ km range and 5-minute refueling, EVs require precise energy budgeting.",
        "A driver stranded due to unexpected range depletion creates safety, logistics, and roadside assistance challenges."
    ]
    for it in items_l:
        p = tf2_l.add_paragraph()
        p.text = "- " + it
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_before = Pt(10)

    # Right card: The Gap
    box2_r = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    box2_r.fill.solid()
    box2_r.fill.fore_color.rgb = COLOR_LIGHT_BG
    box2_r.line.color.rgb = COLOR_BORDER
    tf2_r = box2_r.text_frame
    tf2_r.word_wrap = True
    tf2_r.margin_left = Inches(0.3)
    tf2_r.margin_right = Inches(0.3)
    tf2_r.margin_top = Inches(0.3)

    p = tf2_r.paragraphs[0]
    p.text = "Limitations of Existing Systems"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    items_r = [
        "Official Standard Ratings (WLTP / EPA): Tested under standardized laboratory temperature cycles; do not reflect highway or winter real-world driving.",
        "Dashboard In-Vehicle Estimators: Heavily rely on trailing moving averages of the last 10 to 50 km driven.",
        "Failure Under Volatile Conditions: Abrupt speed increases, cold snaps, or cabin heating cause sudden range drop-offs of 20% to 40%.",
        "Project Objective: Develop an adaptive machine learning model that maps multi-dimensional telemetry directly to continuous remaining range."
    ]
    for it in items_r:
        p = tf2_r.add_paragraph()
        p.text = "- " + it
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 3: 2. Problem Statement
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_slide_layout)
    add_header(s3, "2. Problem Statement", "Technical Challenges")

    col_w = Inches(3.64)
    # Card 1: Aerodynamics & Speed
    c1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), col_w, Inches(5.2))
    c1.fill.solid()
    c1.fill.fore_color.rgb = COLOR_LIGHT_BG
    c1.line.color.rgb = COLOR_BORDER
    tf = c1.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "Speed & Aerodynamic Drag"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    pts1 = [
        "Aerodynamic power loss scales with the cube of vehicle velocity (v^3).",
        "Driving at 110 km/h versus 50 km/h requires disproportionately higher power from the battery pack.",
        "Highway travel severely compresses driving range compared to urban regenerative cycles."
    ]
    for pt in pts1:
        p = tf.add_paragraph()
        p.text = "- " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_before = Pt(10)

    # Card 2: Thermal Kinetics
    c2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.84), Inches(1.6), col_w, Inches(5.2))
    c2.fill.solid()
    c2.fill.fore_color.rgb = COLOR_LIGHT_BG
    c2.line.color.rgb = COLOR_BORDER
    tf = c2.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "Sub-Zero Thermal Impact"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    pts2 = [
        "Li-ion electrolyte viscosity increases and internal resistance rises in cold weather (< 15 deg C).",
        "Available discharge capacity decreases up to 35% in freezing temperatures.",
        "Positive Temperature Coefficient (PTC) cabin heaters consume 1.0 to 4.5 kW directly from the traction pack.",
        "Unlike ICE cars, EVs have no waste engine heat to heat the cabin."
    ]
    for pt in pts2:
        p = tf.add_paragraph()
        p.text = "- " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_before = Pt(10)

    # Card 3: Formulation
    c3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.88), Inches(1.6), col_w, Inches(5.2))
    c3.fill.solid()
    c3.fill.fore_color.rgb = COLOR_LIGHT_BG
    c3.line.color.rgb = COLOR_BORDER
    tf = c3.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "Mathematical Formulation"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    pts3 = [
        "Goal: Construct a multivariate regression function f(X) -> Y.",
        "Input Vector X in R^5:",
        "  1. Speed (km/h)",
        "  2. Ambient Temperature (deg C)",
        "  3. State of Charge (SoC %)",
        "  4. Road Slope (%)",
        "  5. Cabin HVAC Load (kW)",
        "Target Output Y:",
        "  Continuous Remaining Range (km)."
    ]
    for pt in pts3:
        p = tf.add_paragraph()
        p.text = pt
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 4: 3. Literature Survey
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_slide_layout)
    add_header(s4, "3. Literature Survey", "Existing Research & Methodology Comparison")

    table_shape4 = s4.shapes.add_table(4, 4, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tbl4 = table_shape4.table
    tbl4.columns[0].width = Inches(2.2)
    tbl4.columns[1].width = Inches(3.3)
    tbl4.columns[2].width = Inches(3.2)
    tbl4.columns[3].width = Inches(3.0)

    headers4 = ["Approach", "Core Methodology", "Key Advantages", "Critical Limitations"]
    for i, h in enumerate(headers4):
        cell = tbl4.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

    rows4 = [
        ["Physics-Based Models", "Mechanical differential equations (F_tractive = F_aero + F_roll + F_slope).", "Rooted in physical principles; high transparency.", "Requires complex calibration of drag area, friction, and mass; fails to generalize."],
        ["Statistical Moving Averages", "Trailing window average of Wh/km over preceding 10-50 km.", "Computationally trivial; easy to implement on low-cost ECUs.", "Inherently reactive; introduces severe lag during speed or slope transitions."],
        ["Data-Driven Machine Learning", "Supervised regression models trained on real vehicle telemetry (e.g., VED dataset).", "Captures non-linear coupled interactions (HVAC + speed + cold); highly adaptive.", "Requires empirical comparison to establish the optimal regression family (Trees vs. SVR)."]
    ]

    for row_idx, row_data in enumerate(rows4, start=1):
        for col_idx, val in enumerate(row_data):
            cell = tbl4.cell(row_idx, col_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_WHITE if row_idx % 2 == 1 else COLOR_LIGHT_BG
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_DARK_TEXT

    # -------------------------------------------------------------
    # SLIDE 5: 4. Proposed Solution and Methodology
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_slide_layout)
    add_header(s5, "4. Proposed Solution and Methodology", "Architectural Framework")

    # Workflow steps
    steps = [
        ("Step 1: Data Ingestion", "Trip telemetry structured according to the University of Michigan Vehicle Energy Dataset (VED) schema."),
        ("Step 2: Preprocessing", "Feature normalization through StandardScaler within an isolated scikit-learn Pipeline to prevent data leakage."),
        ("Step 3: Multi-Model Training", "Comprehensive training of Random Forest, Gradient Boosting, and Support Vector Regressor (SVR)."),
        ("Step 4: Benchmarking", "Evaluation across MAE, RMSE, and R2 to determine the superior mathematical architecture."),
        ("Step 5: Model Deployment", "Serialization of champion SVR model (.joblib) and integration with an interactive Streamlit UI.")
    ]

    for idx, (title_st, desc_st) in enumerate(steps):
        y_pos = Inches(1.6 + idx * 1.05)
        # Indicator badge
        b = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), y_pos, Inches(3.2), Inches(0.85))
        b.fill.solid()
        b.fill.fore_color.rgb = COLOR_PRIMARY
        b.line.fill.background()
        tf = b.text_frame
        p = tf.paragraphs[0]
        p.text = title_st
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

        # Description box
        db = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.1), y_pos, Inches(8.4), Inches(0.85))
        db.fill.solid()
        db.fill.fore_color.rgb = COLOR_LIGHT_BG
        db.line.color.rgb = COLOR_BORDER
        tf = db.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = desc_st
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_DARK_TEXT

    # -------------------------------------------------------------
    # SLIDE 6: 5. Solution Development
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_slide_layout)
    add_header(s6, "5. Solution Development", "Implementation and Model Configuration")

    # Box 1: Dataset & Parameters
    b1 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    b1.fill.solid()
    b1.fill.fore_color.rgb = COLOR_LIGHT_BG
    b1.line.color.rgb = COLOR_BORDER
    tf = b1.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "Dataset & Physical Parameters"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    items_dev1 = [
        "Vehicle Profile: 2013 Nissan Leaf BEV (24 kWh nominal pack, 21.5 kWh usable capacity).",
        "Volume: 5,000 driving trip records spanning urban, suburban, and highway regimes.",
        "Split: 80% Training (4,000 samples) and 20% Unseen Testing (1,000 samples).",
        "Key Physics Encoded:",
        "  - Base consumption: 140 Wh/km at 50 km/h.",
        "  - Drag power penalty: ((Speed / 50)^1.8) * 35 Wh/km.",
        "  - Road slope resistance: Slope * 15 Wh/km.",
        "  - Auxiliary load: (HVAC kW * 1000) / Speed.",
        "  - Low-temp chemical degradation below 15 deg C."
    ]
    for it in items_dev1:
        p = tf.add_paragraph()
        p.text = it if it.startswith("  -") else "- " + it
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_before = Pt(4)

    # Box 2: Candidate Algorithms
    b2 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    b2.fill.solid()
    b2.fill.fore_color.rgb = COLOR_LIGHT_BG
    b2.line.color.rgb = COLOR_BORDER
    tf = b2.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "Algorithms & Hyperparameters"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    items_dev2 = [
        "Random Forest Regressor:",
        "  - 150 estimators, max depth: 12, min samples split: 4.",
        "  - Ensemble of de-correlated decision trees with bagging.",
        "Gradient Boosting Regressor:",
        "  - 180 estimators, learning rate: 0.08, max depth: 5.",
        "  - Sequential boosting minimizing squared residual errors.",
        "Support Vector Regressor (SVR):",
        "  - Radial Basis Function (RBF) non-linear kernel.",
        "  - Regularization parameter C: 100.0, Epsilon margin: 1.0.",
        "  - Coupled with StandardScaler in a scikit-learn Pipeline.",
        "Deployment: Saved as models/svr_pipeline.joblib."
    ]
    for it in items_dev2:
        p = tf.add_paragraph()
        p.text = it
        p.font.size = Pt(12)
        p.font.bold = True if it.endswith(":") else False
        p.font.color.rgb = COLOR_PRIMARY if it.endswith(":") else COLOR_DARK_TEXT
        p.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 7: 6. Scalability & Performance Analysis - Scoreboard
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_slide_layout)
    add_header(s7, "6. Scalability & Performance Analysis", "Empirical Evaluation on 1,000 Unseen Test Trips")

    # Table
    t_shape7 = s7.shapes.add_table(4, 5, Inches(0.8), Inches(1.6), Inches(11.7), Inches(2.2))
    tbl7 = t_shape7.table
    tbl7.columns[0].width = Inches(3.5)
    tbl7.columns[1].width = Inches(2.0)
    tbl7.columns[2].width = Inches(2.0)
    tbl7.columns[3].width = Inches(2.0)
    tbl7.columns[4].width = Inches(2.2)

    h7 = ["Model Architecture", "MAE (km)", "RMSE (km)", "R2 Score", "Outcome"]
    for i, h in enumerate(h7):
        cell = tbl7.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

    r7 = [
        ["Support Vector Regressor (SVR)", "1.30 km", "1.92 km", "0.9953", "Selected Champion"],
        ["Gradient Boosting Regressor", "1.88 km", "2.69 km", "0.9909", "Evaluated Baseline"],
        ["Random Forest Regressor", "2.52 km", "3.61 km", "0.9835", "Evaluated Baseline"]
    ]
    for r_idx, r_vals in enumerate(r7, start=1):
        for c_idx, val in enumerate(r_vals):
            cell = tbl7.cell(r_idx, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(239, 246, 255) if r_idx == 1 else COLOR_WHITE
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(12)
            p.font.bold = True if (r_idx == 1 or c_idx == 0) else False
            p.font.color.rgb = COLOR_ACCENT if r_idx == 1 else COLOR_DARK_TEXT

    # Lower analysis cards
    b7_1 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.2), Inches(5.6), Inches(2.7))
    b7_1.fill.solid()
    b7_1.fill.fore_color.rgb = COLOR_LIGHT_BG
    b7_1.line.color.rgb = COLOR_BORDER
    tf = b7_1.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "Why SVR Outperformed Decision Trees"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    reasons = [
        "Continuous Manifold: Battery consumption follows continuous physical differential curves (aerodynamic cubic drag, smooth thermal drop).",
        "Step-wise vs. Smooth: Decision trees use orthogonal axis-aligned cuts, producing piece-wise constant step approximations.",
        "RBF Kernel Advantage: Maps inputs to an infinite-dimensional Hilbert space, accurately fitting smooth non-linear curves."
    ]
    for r in reasons:
        p = tf.add_paragraph()
        p.text = "- " + r
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_before = Pt(4)

    b7_2 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.2), Inches(5.7), Inches(2.7))
    b7_2.fill.solid()
    b7_2.fill.fore_color.rgb = COLOR_LIGHT_BG
    b7_2.line.color.rgb = COLOR_BORDER
    tf = b7_2.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "Feature Importance & Sensitivity"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    fi = [
        "1. State of Charge (SoC %): ~59% (primary energy reserve).",
        "2. Ambient Temperature (deg C): ~15.5% (chemical impedance).",
        "3. Vehicle Speed (km/h): ~10.5% (aerodynamic drag load).",
        "4. Road Slope (%): ~8.2% (potential energy climbing load).",
        "5. Cabin HVAC Load (kW): ~7.0% (parasitic traction pack draw)."
    ]
    for f in fi:
        p = tf.add_paragraph()
        p.text = f
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_before = Pt(3)

    # -------------------------------------------------------------
    # SLIDE 8: 6. Scalability & Performance Analysis - Visual & Validation
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_slide_layout)
    add_header(s8, "6. Scalability & Performance Analysis", "Visual Verification and Edge-Case Validation")

    img_path = os.path.join(assets_dir, "ev_model_comparison.png")
    if os.path.exists(img_path):
        s8.shapes.add_picture(img_path, Inches(0.8), Inches(1.5), width=Inches(6.8))

    # Right side: Edge-case tests & Scalability specs
    edge_box = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.5), Inches(4.7), Inches(5.4))
    edge_box.fill.solid()
    edge_box.fill.fore_color.rgb = COLOR_LIGHT_BG
    edge_box.line.color.rgb = COLOR_BORDER
    tf = edge_box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.25)

    p = tf.paragraphs[0]
    p.text = "Real-World Edge Case Validation"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    cases = [
        "Scenario A: Freezing Winter Highway",
        "  - Inputs: -5 deg C, 100 km/h, 80% SoC, 4.0 kW Heater",
        "  - Predicted Range: 39.8 km (62% range penalty)",
        "Scenario B: Mild Spring City Commute",
        "  - Inputs: 20 deg C, 40 km/h, 80% SoC, HVAC Off",
        "  - Predicted Range: 106.3 km (Optimal operating zone)",
        "Scenario C: Hot Summer Incline",
        "  - Inputs: 34 deg C, 60 km/h, 50% SoC, +4% Grade",
        "  - Predicted Range: 37.1 km"
    ]
    for c in cases:
        p = tf.add_paragraph()
        p.text = c
        p.font.size = Pt(11)
        p.font.bold = True if c.startswith("Scenario") else False
        p.font.color.rgb = COLOR_ACCENT if c.startswith("Scenario") else COLOR_DARK_TEXT
        p.space_before = Pt(3)

    p = tf.add_paragraph()
    p.text = "System Scalability Metrics"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_before = Pt(12)

    scal = [
        "- Model Size: < 2 MB (embeddable on vehicle ECU).",
        "- Inference Latency: < 2 milliseconds per sample.",
        "- Deployment: Streamlit web dashboard + standalone CLI."
    ]
    for s in scal:
        p = tf.add_paragraph()
        p.text = s
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_before = Pt(3)

    # -------------------------------------------------------------
    # SLIDE 9: 7. Conclusion
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_slide_layout)
    add_header(s9, "7. Conclusion", "Summary and Future Directions")

    # Left: Conclusion
    c_l = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    c_l.fill.solid()
    c_l.fill.fore_color.rgb = COLOR_LIGHT_BG
    c_l.line.color.rgb = COLOR_BORDER
    tf = c_l.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "Project Summary & Outcomes"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    c_items = [
        "Successfully developed an end-to-end physics-grounded machine learning regression framework for EV driving range prediction.",
        "Evaluated three distinct model families on 1,000 unseen test trips; Support Vector Regression (SVR) demonstrated the highest precision (MAE: 1.30 km, R2: 0.9953).",
        "Demonstrated that ambient temperature and vehicle speed are the primary determinants of non-tractive energy loss.",
        "Successfully deployed the champion SVR pipeline into a real-time interactive Streamlit web dashboard and standalone CLI inference tool."
    ]
    for it in c_items:
        p = tf.add_paragraph()
        p.text = "- " + it
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_before = Pt(10)

    # Right: Future Scope
    c_r = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    c_r.fill.solid()
    c_r.fill.fore_color.rgb = COLOR_LIGHT_BG
    c_r.line.color.rgb = COLOR_BORDER
    tf = c_r.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "Future Scope & Enhancements"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    f_items = [
        "Live In-Vehicle Telematics: Direct integration with OBD-II / CAN-bus hardware for real-time live-streaming range updates.",
        "Battery Degradation Modeling: Incorporate State of Health (SoH) and multi-year battery cell degradation tracking.",
        "Route-Ahead Navigation Integration: Interface with mapping and elevation APIs to predict range for planned future routes before departure.",
        "Driver Behavioral Profiling: Incorporate aggressive versus eco-driving habits into real-time dynamic range scaling."
    ]
    for it in f_items:
        p = tf.add_paragraph()
        p.text = "- " + it
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_before = Pt(10)

    # Save
    prs.save(output_path)
    print(f"Presentation successfully saved to: {output_path}")

if __name__ == "__main__":
    out_file = "/Users/airlord/Desktop/COA/EV_Range_Prediction_Presentation.pptx"
    assets = "/Users/airlord/Desktop/COA/assets"
    build_presentation(out_file, assets)
